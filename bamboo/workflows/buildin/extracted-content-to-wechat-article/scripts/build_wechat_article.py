#!/usr/bin/env python3
"""Build a WeChat article package from extracted content and a Markdown draft."""

from __future__ import annotations

import argparse
import json
import os
import re
import shlex
import sys
import tempfile
import urllib.error
import urllib.request
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path

from PIL import Image

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor
except Exception as exc:  # pragma: no cover - dependency error path
    raise SystemExit(f"python-docx is required: {exc}") from exc


IMAGE_RE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")
LINK_RE = re.compile(r"(?<!!)\[([^\]]+)\]\((https?://[^)]+)\)")
BLOCK_FORMULA_RE = re.compile(r"^\s*\$\$\s*$")


@dataclass(frozen=True)
class Asset:
    kind: str
    path: Path
    rel_path: str
    source_id: str
    source_dir: Path
    page: int | None = None
    caption: str = ""
    width: int = 0
    height: int = 0

    @property
    def key(self) -> str:
        return f"{self.source_id}/{self.rel_path}".replace("\\", "/")

    @property
    def important(self) -> bool:
        if self.kind == "table":
            return self.width >= 500 and self.height >= 120
        return self.width >= 350 and self.height >= 140


@dataclass(frozen=True)
class ImageUse:
    alt: str
    raw_path: str
    resolved_path: Path | None
    key: str


@dataclass(frozen=True)
class LinkStatus:
    label: str
    url: str
    status: str
    detail: str


@dataclass
class BuildStats:
    paragraph_count: int = 0
    image_count: int = 0
    formula_count: int = 0
    formula_render_failures: list[str] | None = None

    def __post_init__(self) -> None:
        if self.formula_render_failures is None:
            self.formula_render_failures = []


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv if argv is not None else sys.argv[1:])
    content_dirs = [path.expanduser().resolve() for path in args.content_dirs]
    article_md = args.article_md.expanduser().resolve()
    output_dir = args.output_dir.expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    for content_dir in content_dirs:
        if not content_dir.is_dir():
            raise SystemExit(f"content dir not found: {content_dir}")
    if not article_md.is_file():
        raise SystemExit(f"article Markdown not found: {article_md}")

    sources = [Source(id=source_id(index, path), path=path) for index, path in enumerate(content_dirs, start=1)]
    assets = discover_assets(sources)
    markdown = article_md.read_text(encoding="utf-8")
    title = args.title or first_heading(markdown) or article_md.stem
    image_uses = find_image_uses(markdown, article_md.parent, sources)
    links = find_links(markdown)
    link_statuses = verify_links(links) if args.verify_links else [
        LinkStatus(label=label, url=url, status="not_checked", detail="run with --verify-links to verify")
        for label, url in links
    ]

    package_assets_dir = output_dir / "assets"
    package_assets_dir.mkdir(exist_ok=True)
    formula_dir = output_dir / "formulas"
    formula_dir.mkdir(exist_ok=True)

    normalized_md = normalize_markdown(markdown, image_uses, output_dir, package_assets_dir)
    normalized_md_path = output_dir / f"{slugify(title)}.normalized.md"
    normalized_md_path.write_text(normalized_md, encoding="utf-8")

    cover_path = args.cover.expanduser().resolve() if args.cover else None
    docx_path = output_dir / f"{slugify(title)}.docx"
    stats = build_docx(
        markdown=normalized_md,
        title=title,
        docx_path=docx_path,
        base_dir=output_dir,
        formula_dir=formula_dir,
        cover_path=cover_path,
    )

    used_keys = {use.key for use in image_uses if use.key}
    manifest_path = output_dir / "asset_manifest.md"
    manifest_path.write_text(render_asset_manifest(assets, used_keys), encoding="utf-8")

    link_report_path = output_dir / "link_report.md"
    link_report_path.write_text(render_link_report(link_statuses), encoding="utf-8")

    quality_report_path = output_dir / "quality_report.md"
    quality_report_path.write_text(
        render_quality_report(
            docx_path=docx_path,
            normalized_md_path=normalized_md_path,
            assets=assets,
            used_keys=used_keys,
            image_uses=image_uses,
            link_statuses=link_statuses,
            stats=stats,
        ),
        encoding="utf-8",
    )

    print(f"DOCX: {docx_path}")
    print(f"Markdown: {normalized_md_path}")
    print(f"AssetManifest: {manifest_path}")
    print(f"LinkReport: {link_report_path}")
    print(f"QualityReport: {quality_report_path}")
    print(f"FormulaAssets: {formula_dir}")
    return 0


def parse_args(argv: list[str]) -> argparse.Namespace:
    tokens = shlex.split(argv[0]) if len(argv) == 1 else argv
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "content_dirs",
        type=Path,
        nargs="+",
        help="One or more extracted content directories. The final two positional arguments are article_md and output_dir.",
    )
    parser.add_argument("article_md", type=Path)
    parser.add_argument("output_dir", type=Path)
    parser.add_argument("--title", default="")
    parser.add_argument("--cover", type=Path)
    parser.add_argument("--no-cover", action="store_true", help="Allow building without a cover image.")
    parser.add_argument("--verify-links", action="store_true")
    args = parser.parse_args(tokens)
    if args.cover and args.no_cover:
        parser.error("--cover and --no-cover cannot be used together")
    if not args.cover and not args.no_cover:
        parser.error("a cover image is required; pass --cover /path/to/cover.png or explicit --no-cover")
    return args


@dataclass(frozen=True)
class Source:
    id: str
    path: Path


def source_id(index: int, path: Path) -> str:
    return f"source-{index:02d}-{slugify(path.name)[:32]}"


def discover_assets(sources: list[Source]) -> list[Asset]:
    by_rel: dict[str, Asset] = {}
    for source in sources:
        content_dir = source.path
        document_json = content_dir / "document.json"
        if document_json.is_file():
            data = json.loads(document_json.read_text(encoding="utf-8"))
            for item in data.get("content", []):
                kind = item.get("type")
                asset_path = item.get("asset_path")
                if kind not in {"figure", "table"} or not asset_path:
                    continue
                path = (content_dir / asset_path).resolve()
                asset = make_asset(
                    kind,
                    path,
                    rel_to(path, content_dir),
                    source.id,
                    content_dir,
                    item.get("page"),
                    item.get("caption") or "",
                )
                if asset:
                    by_rel[asset.key] = asset

        asset_dirs = (
            ("figure", "figures"),
            ("table", "tables"),
            ("image", "images"),
            ("image", "frames"),
            ("image", "screenshots"),
        )
        for kind, dirname in asset_dirs:
            for path in sorted((content_dir / "assets" / dirname).glob("*")):
                if not is_supported_image(path):
                    continue
                key = f"{source.id}/{rel_to(path, content_dir)}".replace("\\", "/")
                if key not in by_rel:
                    asset = make_asset(kind, path.resolve(), rel_to(path, content_dir), source.id, content_dir, None, "")
                    if asset:
                        by_rel[asset.key] = asset
        for path in sorted((content_dir / "assets").glob("*")):
            if not is_supported_image(path):
                continue
            key = f"{source.id}/{rel_to(path, content_dir)}".replace("\\", "/")
            if key not in by_rel:
                asset = make_asset("image", path.resolve(), rel_to(path, content_dir), source.id, content_dir, None, "")
                if asset:
                    by_rel[asset.key] = asset
    return sorted(by_rel.values(), key=lambda item: (item.source_id, item.kind, item.page or 9999, item.rel_path))


def is_supported_image(path: Path) -> bool:
    return path.is_file() and path.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp"}


def make_asset(
    kind: str,
    path: Path,
    rel_path: str,
    source_id_value: str,
    source_dir: Path,
    page: int | None,
    caption: str,
) -> Asset | None:
    if not path.is_file():
        return None
    try:
        with Image.open(path) as img:
            width, height = img.size
    except Exception:
        width, height = 0, 0
    return Asset(
        kind=kind,
        path=path,
        rel_path=rel_path,
        source_id=source_id_value,
        source_dir=source_dir,
        page=page,
        caption=caption,
        width=width,
        height=height,
    )


def find_image_uses(markdown: str, article_dir: Path, sources: list[Source]) -> list[ImageUse]:
    uses: list[ImageUse] = []
    for alt, raw in IMAGE_RE.findall(markdown):
        target = raw.strip().split()[0]
        resolved, source = resolve_image_path(target, article_dir, sources)
        key = f"{source.id}/{rel_to(resolved, source.path)}" if resolved and source else target
        uses.append(ImageUse(alt=alt.strip(), raw_path=target, resolved_path=resolved, key=key.replace("\\", "/")))
    return uses


def resolve_image_path(raw: str, article_dir: Path, sources: list[Source]) -> tuple[Path | None, Source | None]:
    path = Path(raw).expanduser()
    if path.is_absolute():
        candidates: list[tuple[Path, Source | None]] = [(path, source_for_absolute(path, sources))]
    else:
        candidates = [(article_dir / path, None)]
        candidates.extend((source.path / path, source) for source in sources)
    for candidate in candidates:
        candidate_path, source = candidate
        resolved = candidate_path.resolve()
        if resolved.is_file():
            return resolved, source or source_for_absolute(resolved, sources)
    return None, None


def source_for_absolute(path: Path, sources: list[Source]) -> Source | None:
    resolved = path.expanduser().resolve()
    for source in sources:
        if is_relative_to(resolved, source.path):
            return source
    return None


def normalize_markdown(markdown: str, image_uses: list[ImageUse], output_dir: Path, package_assets_dir: Path) -> str:
    replacements: dict[str, str] = {}
    for use in image_uses:
        if not use.resolved_path:
            continue
        dest = unique_asset_destination(package_assets_dir, use.resolved_path.name)
        if not dest.exists():
            dest.write_bytes(use.resolved_path.read_bytes())
        replacements[use.raw_path] = rel_to(dest, output_dir)

    def repl(match: re.Match[str]) -> str:
        alt, raw = match.group(1), match.group(2)
        clean = raw.strip().split()[0]
        return f"![{alt}]({replacements.get(clean, raw)})"

    return IMAGE_RE.sub(repl, markdown)


def unique_asset_destination(directory: Path, filename: str) -> Path:
    candidate = directory / filename
    if not candidate.exists():
        return candidate
    stem = candidate.stem
    suffix = candidate.suffix
    index = 2
    while True:
        candidate = directory / f"{stem}-{index}{suffix}"
        if not candidate.exists():
            return candidate
        index += 1


def build_docx(
    *,
    markdown: str,
    title: str,
    docx_path: Path,
    base_dir: Path,
    formula_dir: Path,
    cover_path: Path | None,
) -> BuildStats:
    doc = Document()
    stats = BuildStats()
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)

    if cover_path and cover_path.is_file():
        add_image(doc, cover_path, width=6.0)
        stats.image_count += 1

    in_formula = False
    formula_lines: list[str] = []
    formula_index = 1
    for raw_line in markdown.splitlines():
        line = raw_line.rstrip()
        if BLOCK_FORMULA_RE.match(line):
            if in_formula:
                formula = "\n".join(formula_lines).strip()
                formula_path = formula_dir / f"formula_{formula_index:03d}.png"
                formula_index += 1
                stats.formula_count += 1
                if render_formula_png(formula, formula_path):
                    add_image(doc, formula_path, width=5.7)
                    stats.image_count += 1
                else:
                    stats.formula_render_failures.append(formula)
                    add_paragraph(doc, f"$${formula}$$", italic=True, color=RGBColor(0x66, 0x66, 0x66))
                    stats.paragraph_count += 1
                formula_lines = []
                in_formula = False
            else:
                in_formula = True
                formula_lines = []
            continue
        if in_formula:
            formula_lines.append(line)
            continue

        image_match = IMAGE_RE.search(line)
        if image_match:
            alt, raw = image_match.group(1).strip(), image_match.group(2).strip()
            image_path = (base_dir / raw).resolve() if not Path(raw).is_absolute() else Path(raw).resolve()
            if image_path.is_file():
                add_image(doc, image_path, width=5.8)
                stats.image_count += 1
                if alt:
                    add_caption(doc, alt)
                    stats.paragraph_count += 1
            else:
                add_paragraph(doc, f"[missing image: {raw}] {alt}", italic=True, color=RGBColor(0xAA, 0x33, 0x33))
                stats.paragraph_count += 1
            continue

        if not line.strip():
            continue
        if line.startswith("# "):
            doc.add_heading(strip_markdown(line[2:]), level=0)
        elif line.startswith("## "):
            doc.add_heading(strip_markdown(line[3:]), level=1)
        elif line.startswith("### "):
            doc.add_heading(strip_markdown(line[4:]), level=2)
        elif line.startswith("- "):
            add_paragraph(doc, strip_markdown(line[2:]), style="List Bullet")
            stats.paragraph_count += 1
        elif re.match(r"^\d+\.\s+", line):
            add_paragraph(doc, strip_markdown(re.sub(r"^\d+\.\s+", "", line)), style="List Number")
            stats.paragraph_count += 1
        elif line.startswith("> "):
            add_paragraph(doc, strip_markdown(line[2:]), italic=True, color=RGBColor(0x55, 0x55, 0x55))
            stats.paragraph_count += 1
        else:
            add_paragraph(doc, strip_markdown(line))
            stats.paragraph_count += 1

    doc.save(docx_path)
    return stats


def add_paragraph(doc: Document, text: str, *, style: str | None = None, italic: bool = False, color: RGBColor | None = None) -> None:
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(text)
    run.font.size = Pt(11)
    run.italic = italic
    if color:
        run.font.color.rgb = color


def add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)


def add_image(doc: Document, path: Path, *, width: float) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))


def render_formula_png(formula: str, out_path: Path) -> bool:
    try:
        os.environ.setdefault("MPLCONFIGDIR", tempfile.mkdtemp(prefix="bamboo-mpl-"))
        import matplotlib.pyplot as plt

        text = f"${formula}$"
        fig = plt.figure(figsize=(8, 0.8), dpi=220)
        fig.patch.set_alpha(0)
        fig.text(0.5, 0.5, text, ha="center", va="center", fontsize=18)
        fig.savefig(out_path, bbox_inches="tight", pad_inches=0.15, transparent=True)
        plt.close(fig)
        return True
    except Exception:
        return False


def strip_markdown(text: str) -> str:
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1（\2）", text)
    return text


def first_heading(markdown: str) -> str:
    for line in markdown.splitlines():
        if line.startswith("# "):
            return strip_markdown(line[2:].strip())
    return ""


def find_links(markdown: str) -> list[tuple[str, str]]:
    return [(label.strip(), url.strip()) for label, url in LINK_RE.findall(markdown)]


def verify_links(links: Iterable[tuple[str, str]]) -> list[LinkStatus]:
    statuses: list[LinkStatus] = []
    for label, url in links:
        request = urllib.request.Request(url, method="HEAD", headers={"User-Agent": "BambooLinkChecker/1.0"})
        try:
            with urllib.request.urlopen(request, timeout=10) as response:
                statuses.append(LinkStatus(label, url, "ok", str(response.status)))
        except urllib.error.HTTPError as exc:
            if exc.code in {403, 405}:
                statuses.append(verify_link_get(label, url))
            else:
                statuses.append(LinkStatus(label, url, "failed", f"HTTP {exc.code}"))
        except Exception as exc:
            statuses.append(LinkStatus(label, url, "failed", str(exc)))
    return statuses


def verify_link_get(label: str, url: str) -> LinkStatus:
    request = urllib.request.Request(url, method="GET", headers={"User-Agent": "BambooLinkChecker/1.0"})
    try:
        with urllib.request.urlopen(request, timeout=10) as response:
            status = "ok" if response.status < 400 else "failed"
            return LinkStatus(label, url, status, str(response.status))
    except Exception as exc:
        return LinkStatus(label, url, "failed", str(exc))


def render_asset_manifest(assets: list[Asset], used_keys: set[str]) -> str:
    rows = [
        "# Asset Manifest",
        "",
        "| Used | Important | Source | Type | Page | Size | Path | Caption |",
        "| --- | --- | --- | --- | --- | --- | --- | --- |",
    ]
    for asset in assets:
        rows.append(
            "| {used} | {important} | `{source}` | {kind} | {page} | {size} | `{path}` | {caption} |".format(
                used="yes" if asset.key in used_keys else "no",
                important="yes" if asset.important else "no",
                source=asset.source_id,
                kind=asset.kind,
                page=asset.page or "",
                size=f"{asset.width}x{asset.height}",
                path=asset.rel_path,
                caption=(asset.caption or "").replace("|", "\\|")[:220],
            )
        )
    return "\n".join(rows) + "\n"


def render_link_report(statuses: list[LinkStatus]) -> str:
    rows = [
        "# Link Report",
        "",
        "| Status | Label | URL | Detail |",
        "| --- | --- | --- | --- |",
    ]
    for item in statuses:
        label = item.label.replace("|", "\\|")
        detail = item.detail.replace("|", "\\|")
        rows.append(f"| {item.status} | {label} | {item.url} | {detail} |")
    return "\n".join(rows) + "\n"


def render_quality_report(
    *,
    docx_path: Path,
    normalized_md_path: Path,
    assets: list[Asset],
    used_keys: set[str],
    image_uses: list[ImageUse],
    link_statuses: list[LinkStatus],
    stats: BuildStats,
) -> str:
    important_unused = [asset for asset in assets if asset.important and asset.key not in used_keys]
    missing_images = [use for use in image_uses if use.resolved_path is None]
    failed_links = [item for item in link_statuses if item.status == "failed"]
    unchecked_links = [item for item in link_statuses if item.status == "not_checked"]
    lines = [
        "# Quality Report",
        "",
        f"- DOCX: `{docx_path}`",
        f"- Markdown: `{normalized_md_path}`",
        f"- Paragraphs: {stats.paragraph_count}",
        f"- Embedded images: {stats.image_count}",
        f"- Block formulas: {stats.formula_count}",
        f"- Formula render failures: {len(stats.formula_render_failures or [])}",
        f"- Article image references: {len(image_uses)}",
        f"- Extracted assets: {len(assets)}",
        f"- Important unused assets: {len(important_unused)}",
        f"- Missing image references: {len(missing_images)}",
        f"- Links: {len(link_statuses)}",
        f"- Failed links: {len(failed_links)}",
        f"- Unchecked links: {len(unchecked_links)}",
        "",
    ]
    if important_unused:
        lines.extend(["## Important Unused Assets", ""])
        for asset in important_unused:
            lines.append(f"- `{asset.key}` ({asset.kind}, page {asset.page or '?'}, {asset.width}x{asset.height})")
        lines.append("")
    if missing_images:
        lines.extend(["## Missing Image References", ""])
        for use in missing_images:
            lines.append(f"- `{use.raw_path}` ({use.alt})")
        lines.append("")
    if failed_links:
        lines.extend(["## Failed Links", ""])
        for item in failed_links:
            lines.append(f"- {item.url}: {item.detail}")
        lines.append("")
    if stats.formula_render_failures:
        lines.extend(["## Formula Render Failures", ""])
        for formula in stats.formula_render_failures:
            lines.append("```latex")
            lines.append(formula)
            lines.append("```")
        lines.append("")
    return "\n".join(lines)


def rel_to(path: Path, root: Path) -> str:
    try:
        return str(path.resolve().relative_to(root.resolve()))
    except ValueError:
        return str(path)


def is_relative_to(path: Path, root: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
        return True
    except ValueError:
        return False


def slugify(value: str) -> str:
    cleaned = re.sub(r"[\\/:*?\"<>|]+", "-", value).strip()
    cleaned = re.sub(r"\s+", "-", cleaned)
    return cleaned[:80] or "wechat-article"


if __name__ == "__main__":
    raise SystemExit(main())
